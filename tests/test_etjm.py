"""Tests for the journal maker. Run with:  python -m unittest discover -s tests

Everything here uses made-up names and jobs.
"""

import csv
import io
import json
import os
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))

from etjm_accounts import (  # noqa: E402
    Keys,
    apply_password_changes,
    check_own_password,
    load_state,
    new_state,
    profile_for,
    public_users,
    sync_people,
)
from etjm_core import (  # noqa: E402
    PASSWORD_V1,
    TIMESHEET_V2,
    TimesheetError,
    b64,
    bundle_key,
    clean_name,
    collect_timesheets,
    decrypt_bundle,
    default_username,
    encrypt_bundle,
    encrypt_timesheet,
    generate_password,
    load_settings,
    normalize_password,
    parse_timesheet,
    people,
    read_job_report,
    seal,
    unseal,
    unseal_profile,
    unwrap_user_key,
    user_id,
    user_tag,
    wrap_user_key,
)
from etjm_outputs import build_all, build_journal, journal_csv  # noqa: E402

MASTER = bytes(range(32))


def issue(number, data, updated="2026-10-01T00:00:00Z", **extra):
    """In the journal tests an "issue" is just the saved data (see collected())."""
    return data


def collected(datas):
    out = []
    for data in datas:
        timesheet = parse_timesheet(data)
        timesheet.username = default_username(timesheet.employee)
        out.append(timesheet)
    return out, []


def sheet(employee="Ann Example", month="2026-09", jobs=None, other=None, saved="2026-10-01T00:00:00Z"):
    return {"v": 1, "employee": employee, "month": month, "savedAt": saved,
            "jobs": jobs if jobs is not None else [], "other": other if other is not None else []}


def settings(**changes):
    result = load_settings(None)
    result["employees"] = [{"name": "Ann Example", "rechargeRate": 160}, {"name": "Ben Sample", "rechargeRate": 215}]
    result.update(changes)
    return result


def world():
    """Two engineers and an admin with sign-ins, as the Actions would set them up."""
    keys = Keys(MASTER)
    everyone = [
        {"name": "Ann Example", "username": "ann.example", "timesheets": True, "admin": False},
        {"name": "Ben Sample", "username": "ben.sample", "timesheets": True, "admin": False},
        {"name": "Dee Office", "username": "dee.office", "timesheets": False, "admin": True},
    ]
    state = new_state()
    sync_people(state, keys, everyone)
    return keys, everyone, state


def keys_by_tag(keys, state, names=None):
    names = names or {"ann.example": "Ann Example", "ben.sample": "Ben Sample"}
    return {user_tag(keys.master, u): {"username": u, "name": names[u], "dataKeys": keys.data_keys(u, r)}
            for u, r in state["users"].items() if u in names}


def ts_issue(number, keys, state, username, data, key_index=None, updated="2026-10-01T00:00:00Z", state_="open"):
    record = state["users"][username]
    data_keys = keys.data_keys(username, record)
    index = len(data_keys) - 1 if key_index is None else key_index
    token = encrypt_timesheet(data, data_keys[index], index, username)
    return {"number": number, "state": state_, "updated_at": updated,
            "title": f"Timesheet {data['month']} · {user_tag(keys.master, username)}",
            "body": f"{TIMESHEET_V2}\nEncrypted.\n\n```etjm\n{token}\n```\n"}


class NameTests(unittest.TestCase):
    def test_good_names(self):
        self.assertEqual(clean_name("  Mary-Jane  O’Neil "), "Mary-Jane O’Neil")
        self.assertEqual(clean_name("Tāne Mahuta"), "Tāne Mahuta")

    def test_bad_names(self):
        for bad in ["", "=cmd()", "+SUM(1)", "-1", "@x", "R2D2", "Ann <b>", "x" * 61]:
            with self.subTest(bad=bad), self.assertRaises(TimesheetError):
                clean_name(bad)


class TimesheetTests(unittest.TestCase):
    def test_parse_and_totals(self):
        ts = parse_timesheet(sheet(jobs=[
            {"job": "JOB000001", "location": "SE", "owner": "80", "description": "Tank",
             "hours": {"REVEX": {"1": 2, "2": "1.5"}, "CAPEX": {"3": 0.25}}},
            {"job": "job000001", "hours": {"REVEX": {"1": 1}}},  # same job again: added together
        ], other=[{"label": "Annual Leave", "hours": {"4": 8}}]))
        self.assertEqual(ts.month_key, "2026-09")
        lines = {(line.job, line.task): line for line in ts.jobs}
        self.assertEqual(lines[("JOB000001", "REVEX")].hours, {1: Decimal(3), 2: Decimal("1.5")})
        self.assertEqual(ts.job_total, Decimal("4.75"))
        self.assertEqual(ts.other_total, Decimal(8))
        self.assertEqual(ts.day_totals()[1], Decimal(3))

    def test_rejects_bad_data(self):
        cases = {
            "too many hours": {"REVEX": {"1": 25}},
            "negative": {"REVEX": {"1": -1}},
            "text": {"REVEX": {"1": "lots"}},
            "three decimals": {"REVEX": {"1": 1.125}},
            "day 31 in September": {"REVEX": {"31": 1}},
            "not CAPEX/REVEX": {"OPEX": {"1": 1}},
        }
        for name, hours in cases.items():
            with self.subTest(name), self.assertRaises(TimesheetError):
                parse_timesheet(sheet(jobs=[{"job": "JOB000001", "hours": hours}]))
        with self.assertRaises(TimesheetError):
            parse_timesheet(sheet(jobs=[{"job": "J1234", "hours": {}}]))
        with self.assertRaises(TimesheetError):
            parse_timesheet(sheet(month="2026-13"))


class CryptoTests(unittest.TestCase):
    def test_sealed_data_is_bound_to_its_purpose(self):
        key = bytes(32)
        token = seal(key, b"secret", "etjm-ts|ann.example|2026-09")
        self.assertEqual(unseal(key, token, "etjm-ts|ann.example|2026-09"), b"secret")
        for wrong in ("etjm-ts|ben.sample|2026-09", "etjm-ts|ann.example|2026-10"):
            with self.assertRaises(ValueError):
                unseal(key, token, wrong)
        with self.assertRaises(ValueError):
            unseal(bytes(range(32)), token, "etjm-ts|ann.example|2026-09")

    def test_password_unlocks_the_user_key(self):
        user_key = os.urandom(32)
        wrap = wrap_user_key(user_key, "abcd-efgh-2345", "ann.example", iterations=1000)
        self.assertEqual(unwrap_user_key(wrap, "ABCD EFGH 2345", "ann.example"), user_key)
        with self.assertRaises(ValueError):
            unwrap_user_key(wrap, "ABCD-EFGH-2346", "ann.example")
        with self.assertRaises(ValueError):
            unwrap_user_key(wrap, "abcd-efgh-2345", "ben.sample")

    def test_bundle_round_trip_and_epochs(self):
        blob = encrypt_bundle({"jobs": [1]}, bundle_key(MASTER, 3), 3)
        self.assertNotIn("jobs", json.dumps(blob))
        self.assertEqual(decrypt_bundle(blob, bundle_key(MASTER, 3)), {"jobs": [1]})
        with self.assertRaises(ValueError):
            decrypt_bundle(blob, bundle_key(MASTER, 2))

    def test_generated_passwords_and_ids(self):
        self.assertRegex(generate_password(), r"^[A-Z2-9]{4}-[A-Z2-9]{4}-[A-Z2-9]{4}$")
        self.assertRegex(generate_password(4), r"^([A-Z2-9]{4}-){3}[A-Z2-9]{4}$")
        self.assertEqual(user_id(" Ann.Example "), user_id("ann.example"))
        self.assertNotIn("ann", user_id("ann.example"))
        self.assertNotIn("ann", user_tag(MASTER, "ann.example"))


class AccountTests(unittest.TestCase):
    def test_new_people_get_keys_and_issued_passwords(self):
        keys, everyone, state = world()
        ann, dee = state["users"]["ann.example"], state["users"]["dee.office"]
        password = keys.password("ann.example", ann)
        self.assertRegex(password, r"^[A-Z2-9]{4}(-[A-Z2-9]{4}){2}$")
        self.assertRegex(keys.password("dee.office", dee), r"^[A-Z2-9]{4}(-[A-Z2-9]{4}){3}$", "admins get longer passwords")
        self.assertEqual(unwrap_user_key(ann["wrap"], password, "ann.example"), keys.user_key("ann.example", ann))
        self.assertNotIn(password, json.dumps(state), "passwords are never stored in plain text")
        self.assertEqual(state["epoch"], 1)

    def test_reset_and_removal_lock_people_out(self):
        keys, everyone, state = world()
        old_key = keys.user_key("ann.example", state["users"]["ann.example"])
        state["users"]["ann.example"]["reset"] = True
        report = sync_people(state, keys, everyone)
        ann = state["users"]["ann.example"]
        self.assertEqual(report["reset"], ["ann.example"])
        self.assertNotEqual(keys.user_key("ann.example", ann), old_key, "old sessions stop working")
        self.assertEqual(len(ann["dataKeys"]), 2)
        self.assertNotIn("reset", ann)
        self.assertEqual(state["epoch"], 2, "the job-list key changes")
        report = sync_people(state, keys, [p for p in everyone if p["username"] != "ben.sample"])
        self.assertEqual(report["removed"], ["ben.sample"])
        self.assertNotIn("ben.sample", state["users"])
        self.assertEqual(state["epoch"], 3)

    def test_password_change_requests(self):
        keys, everyone, state = world()
        ann = state["users"]["ann.example"]
        user_key = keys.user_key("ann.example", ann)
        tag = user_tag(MASTER, "ann.example")

        def request(number, seq, password, sealing_key=user_key, username="ann.example"):
            wrap = wrap_user_key(user_key, password, "ann.example")
            token = "ETJMP1." + seal(sealing_key, json.dumps({"seq": seq, "wrap": wrap}).encode(), f"etjm-pw|{username}")
            return {"number": number, "state": "open", "title": f"Password change · {tag}",
                    "body": f"{PASSWORD_V1}\n```etjm\n{token}\n```\n"}

        done, notes = apply_password_changes(state, keys, [request(5, 100, "my own long secret")])
        self.assertEqual(done, [5])
        self.assertTrue(ann["own"])
        self.assertNotIn("password", ann, "the issued password is forgotten")
        self.assertEqual(unwrap_user_key(ann["wrap"], "my own long secret", "ann.example"), user_key)
        replay, _ = apply_password_changes(state, keys, [request(6, 100, "another password")])
        self.assertEqual(unwrap_user_key(ann["wrap"], "my own long secret", "ann.example"), user_key, "a replayed request is ignored")
        forged, notes = apply_password_changes(state, keys, [request(7, 200, "hacker password", sealing_key=os.urandom(32))])
        self.assertEqual(forged, [7])
        self.assertEqual(unwrap_user_key(ann["wrap"], "my own long secret", "ann.example"), user_key, "a forged request is ignored")
        self.assertIn("couldn't be verified", notes[0])

    def test_profiles_and_public_file(self):
        keys, everyone, state = world()
        conf = settings(admins=[{"name": "Dee Office"}])
        by_name = {p["username"]: p for p in everyone}
        profiles = {u: profile_for(by_name[u], conf, keys, 1, state["users"][u], "admin-token") for u in by_name}
        ann, dee = profiles["ann.example"], profiles["dee.office"]
        self.assertEqual((ann["rate"], len(ann["dataKeys"])), (160, 1))
        self.assertNotIn("passwordKey", ann)
        self.assertNotIn("adminToken", ann)
        self.assertEqual(dee["adminToken"], "admin-token")
        self.assertIn("passwordKey", dee)
        self.assertNotIn("dataKeys", dee)
        public = public_users(state, keys, profiles, {})
        text = json.dumps(public)
        for secret in ("admin-token", "Ann Example", "ann.example", ann["bundleKey"], keys.password("ann.example", state["users"]["ann.example"])):
            self.assertNotIn(secret, text)
        entry = public["users"][user_id("ann.example")]
        user_key = unwrap_user_key(entry, keys.password("ann.example", state["users"]["ann.example"]), "ann.example")
        self.assertEqual(unseal_profile(user_key, entry["profile"], "ann.example"), ann)
        self.assertEqual(public_users(state, keys, profiles, public["users"]), public, "unchanged entries are kept")
        old_format = {user_id("ann.example"): {"salt": "x", "iv": "y", "key": "z", "profile": {"iv": "a", "ct": "b"}}}
        self.assertEqual(public_users(state, keys, profiles, old_format)["users"].keys(), public["users"].keys(), "old-format entries are replaced")

    def test_old_plain_text_file_is_replaced(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "users.json"
            path.write_text(json.dumps({"passwords": {"ann.example": "PLAIN-TEXT-1"}}))
            state, restarted = load_state(path)
        self.assertTrue(restarted)
        self.assertEqual(state["users"], {})

    def test_own_password_rules(self):
        self.assertIsNone(check_own_password("Correct horse battery", "ann.example"))
        self.assertIn("12", check_own_password("short1", "ann.example"))
        self.assertIsNotNone(check_own_password("aaaaaaaaaaaaaaaa", "ann.example"))
        self.assertIsNotNone(check_own_password("annexample-rules-ok", "ann.example"))


class TimesheetSecurityTests(unittest.TestCase):
    def test_only_the_owner_can_save_their_timesheet(self):
        keys, everyone, state = world()
        tags = keys_by_tag(keys, state)
        good = ts_issue(1, keys, state, "ann.example", sheet(jobs=[{"job": "JOB000010", "hours": {"REVEX": {"1": 2}}}]))
        found, warnings = collect_timesheets([good], tags)
        self.assertEqual([(t.username, t.employee, t.job_total) for t in found], [("ann.example", "Ann Example", 2)])

        # Ben encrypts a timesheet with his own key but titles it as Ann's
        forged = ts_issue(2, keys, state, "ben.sample", sheet(employee="Ann Example", saved="2026-10-05T00:00:00Z",
                                                              jobs=[{"job": "JOB000010", "hours": {"REVEX": {"1": 20}}}]))
        forged["title"] = good["title"]
        # someone edits Ann's issue body
        tampered = dict(good, number=3, body=good["body"].replace("ETJM2.0.", "ETJM2.0.A"))
        found, warnings = collect_timesheets([good, forged, tampered], tags)
        self.assertEqual([t.job_total for t in found], [2])
        self.assertEqual(len(warnings), 2)

    def test_closed_or_replayed_issues_cannot_undo_a_save(self):
        keys, everyone, state = world()
        tags = keys_by_tag(keys, state)
        old = ts_issue(1, keys, state, "ann.example", sheet(saved="2026-10-01T00:00:00Z", jobs=[{"job": "JOB000010", "hours": {"REVEX": {"1": 2}}}]))
        new = ts_issue(1, keys, state, "ann.example", sheet(saved="2026-10-02T00:00:00Z", jobs=[{"job": "JOB000010", "hours": {"REVEX": {"1": 5}}}]),
                       state_="closed")
        mirror_ts, _ = collect_timesheets([new], tags)  # closing doesn't matter
        self.assertEqual(mirror_ts[0].job_total, 5)
        mirror = {("ann.example", "2026-09"): mirror_ts[0]}
        found, _ = collect_timesheets([old], tags, mirror)  # the older copy put back into the issue
        self.assertEqual(found[0].job_total, 5)

    def test_old_keys_and_wrong_months_are_rejected(self):
        keys, everyone, state = world()
        state["users"]["ann.example"]["reset"] = True
        sync_people(state, keys, everyone)
        tags = keys_by_tag(keys, state)
        stale = ts_issue(1, keys, state, "ann.example", sheet(), key_index=0)
        moved = ts_issue(2, keys, state, "ann.example", sheet())
        moved["title"] = moved["title"].replace("2026-09", "2026-10")
        found, warnings = collect_timesheets([stale, moved], tags)
        self.assertEqual(found, [])
        self.assertIn("old key", warnings[0])

    def test_name_comes_from_settings_not_the_timesheet(self):
        keys, everyone, state = world()
        tags = keys_by_tag(keys, state)
        spoof = ts_issue(1, keys, state, "ann.example", sheet(employee="Ben Sample", jobs=[{"job": "JOB000010", "hours": {"REVEX": {"1": 2}}}]))
        found, _ = collect_timesheets([spoof], tags)
        self.assertEqual(found[0].employee, "Ann Example")


class JobReportTests(unittest.TestCase):
    def test_reads_by_header_name(self):
        import openpyxl

        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.append(["Status", "Owner", "Project Task No.", "Budget Cost", "Job Description", "Project No.", "Job Location"])
        ws.append(["Open", 80, "REVEX", 100, "Tank repaint", "JOB000002", "SE"])
        ws.append(["Open", 80, "capex", 100, "Tank repaint", "JOB000002", "SE"])
        ws.append(["Completed", 30, "REVEX", 5, "Old job", "JOB000001", "NE"])
        ws.append(["Open", "99", "REVEX", 5, "Pipe", "JOB000003", "DU"])
        ws.append(["Planning", "99", "CAPEX", 5, "New", "JOB000004", "DU"])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "report.xlsx"
            workbook.save(path)
            jobs, stats = read_job_report(path, ["Open"])
        self.assertEqual([j["no"] for j in jobs], ["JOB000002", "JOB000003"])
        self.assertEqual(jobs[0], {"no": "JOB000002", "desc": "Tank repaint", "loc": "SE", "owner": "80", "tasks": ["CAPEX", "REVEX"]})
        self.assertEqual(stats["both"], 1)
        self.assertNotIn("Budget", json.dumps(jobs))


class JournalTests(unittest.TestCase):
    def test_csv_matches_the_journal_layout(self):
        timesheets, _ = collected([
            issue(1, sheet(employee="Ben Sample", jobs=[
                {"job": "JOB000010", "hours": {"REVEX": {"1": 4, "2": 0.5}}},
                {"job": "JOB000011", "hours": {"CAPEX": {"3": 2}, "REVEX": {}}},
            ])),
            issue(2, sheet(employee="Ann Example", jobs=[{"job": "JOB000012", "hours": {"REVEX": {"5": 3}}}],
                           other=[{"label": "Training", "hours": {"6": 8}}])),
            issue(3, sheet(employee="Cat Nojobs", other=[{"label": "Annual Leave", "hours": {"7": 8}}])),
        ])
        lines, warnings = build_journal(timesheets, settings(), 2026, 9)
        self.assertEqual(warnings, [])
        text = journal_csv(lines, settings()).decode("utf-8")
        self.assertTrue(text.startswith("﻿Employee - ID,Date,GST inclusive ,GST,Excl,GL,Project,"))
        rows = list(csv.reader(io.StringIO(text.lstrip("﻿"))))
        self.assertEqual(rows[1:], [
            ["Ann Example", "30/09/2026", "480.00", "0.00", "480.00", "", "JOB000012", "Ann Example Timewriting Sept 26", "REVEX", "3", "160", "EngtimewriteSept 26"],
            ["Ann Example", "30/09/2026", "-480.00", "0.00", "-480.00", "9011", "", "Ann Example Timewriting Sept 26", "", "", "", "EngtimewriteSept 26"],
            ["Ben Sample", "30/09/2026", "967.50", "0.00", "967.50", "", "JOB000010", "Ben Sample Timewriting Sept 26", "REVEX", "4.5", "215", "EngtimewriteSept 26"],
            ["Ben Sample", "30/09/2026", "430.00", "0.00", "430.00", "", "JOB000011", "Ben Sample Timewriting Sept 26", "CAPEX", "2", "215", "EngtimewriteSept 26"],
            ["Ben Sample", "30/09/2026", "-1397.50", "0.00", "-1397.50", "9011", "", "Ben Sample Timewriting Sept 26", "", "", "", "EngtimewriteSept 26"],
        ])
        self.assertIn("\r\n", text)

    def test_month_files_and_checks(self):
        import openpyxl

        data = [issue(1, sheet(employee="Ann Example", jobs=[
            {"job": "JOB000012", "hours": {"REVEX": {"5": 3}}},
            {"job": "JOB000013", "hours": {"REVEX": {"6": 2}}},
        ], other=[{"label": "Training", "hours": {"5": 22}}]))]
        jobs_index = {"JOB000012": {"no": "JOB000012", "desc": "Tank", "loc": "SE", "owner": "80", "tasks": ["CAPEX"]}}
        timesheets, _ = collected(data)
        files, checks = build_all(timesheets, settings(), jobs_index)
        names = sorted(files)
        self.assertIn("output/2026-09/Timewriting Journal - Sep 2026.csv", names)
        self.assertIn("output/2026-09/Timewriting Journal - Sep 2026.xlsx", names)
        self.assertIn("output/2026-09/Engineer timewriting - Ann Example - September 2026.xlsx", names)
        self.assertIn("timesheets/2026-09/ann.example.json", names)
        month_checks = " ".join(checks["2026-09"])
        self.assertIn("JOB000012 is open for CAPEX, not REVEX", month_checks)
        self.assertIn("JOB000013 REVEX (2 h) is not in the current list", month_checks)
        self.assertIn("25 hours entered", month_checks)

        again, _ = build_all(timesheets, settings(), jobs_index)
        self.assertEqual(files, again, "rebuilding the same data must give identical files")

        with tempfile.TemporaryDirectory() as tmp:
            journal = Path(tmp) / "j.xlsx"
            journal.write_bytes(files["output/2026-09/Timewriting Journal - Sep 2026.xlsx"])
            ws = openpyxl.load_workbook(journal)["Journal"]
            self.assertEqual(ws["C2"].value, "=K2*J2")
            self.assertEqual(ws["C4"].value, "=-SUM(C2:C3)")
            self.assertEqual(ws["F4"].value, 9011)
            self.assertEqual(ws["G1"].fill.fgColor.rgb, "FFFFFF00")
            engineer = Path(tmp) / "e.xlsx"
            engineer.write_bytes(files["output/2026-09/Engineer timewriting - Ann Example - September 2026.xlsx"])
            ws = openpyxl.load_workbook(engineer, data_only=True).active
            self.assertEqual(ws.title, "Timewriting 1-For Engineer")
            self.assertEqual(ws["D6"].value, "Training")
            self.assertEqual(ws["A7"].value, "JOB000012")
            self.assertEqual(ws["F7"].value, "REVEX not open")
            self.assertEqual(ws["G4"].value, 27)


class CommitAndPushTests(unittest.TestCase):
    def test_missing_folders_then_add_then_remove(self):
        import subprocess

        from etjm_github import commit_and_push

        def run(*args, cwd):
            subprocess.run(args, cwd=cwd, check=True, capture_output=True)

        with tempfile.TemporaryDirectory() as tmp:
            remote, work = Path(tmp) / "remote.git", Path(tmp) / "work"
            run("git", "init", "-q", "--bare", "-b", "main", str(remote), cwd=tmp)
            run("git", "clone", "-q", str(remote), str(work), cwd=tmp)
            run("git", "config", "user.name", "Test", cwd=work)
            run("git", "config", "user.email", "test@example.com", cwd=work)
            (work / "settings.json").write_text("{}")
            run("git", "add", ".", cwd=work)
            run("git", "commit", "-q", "-m", "start", cwd=work)
            run("git", "push", "-q", "origin", "main", cwd=work)

            nothing = lambda: None  # noqa: E731 - no timesheets yet, so no output/ folder
            self.assertEqual(commit_and_push(work, ["output", "timesheets"], "build", nothing), (False, None))

            def make():
                (work / "output").mkdir(exist_ok=True)
                (work / "output" / "a.csv").write_text("x")
            self.assertTrue(commit_and_push(work, ["output", "timesheets"], "build", make)[0])

            def remove():
                import shutil
                shutil.rmtree(work / "output")
            self.assertTrue(commit_and_push(work, ["output", "timesheets"], "build", remove)[0])
            listing = subprocess.run(["git", "ls-tree", "-r", "--name-only", "origin/main"], cwd=work,
                                     capture_output=True, text=True, check=True).stdout.split()
            self.assertEqual(listing, ["settings.json"])


class RatesAndClosedMonthsTests(unittest.TestCase):
    def test_no_rate_means_zero_and_a_check(self):
        timesheets, _ = collected([issue(1, sheet(employee="Dan Norate", jobs=[{"job": "JOB000010", "hours": {"REVEX": {"1": 2}}}]))])
        lines, warnings = build_journal(timesheets, settings(), 2026, 9)
        self.assertEqual(lines[0].amount, 0)
        self.assertTrue(any("no rechargeRate" in w for w in warnings))

    def test_former_engineers_keep_their_rate_but_cannot_sign_in(self):
        conf = settings()
        conf["employees"] = [{"name": "Ann Example", "username": "ann.example", "rechargeRate": 160, "active": True},
                             {"name": "Ben Sample", "username": "ben", "rechargeRate": 215, "active": False}]
        self.assertEqual([p["username"] for p in people(conf)], ["ann.example"])
        timesheets, _ = collected([issue(1, sheet(employee="Ben Sample", jobs=[{"job": "JOB000010", "hours": {"REVEX": {"1": 2}}}]))])
        lines, _ = build_journal(timesheets, conf, 2026, 9)
        self.assertEqual(lines[0].amount, Decimal("430.00"))

    def test_closed_months_are_frozen(self):
        from build_outputs import write_tree
        from etjm_vault import Vault

        timesheets, _ = collected([
            issue(1, sheet(month="2026-08", jobs=[{"job": "JOB000010", "hours": {"REVEX": {"1": 2}}}])),
            issue(2, sheet(month="2026-09", jobs=[{"job": "JOB000010", "hours": {"REVEX": {"1": 3}}}])),
        ])
        conf = settings(closedMonths=["2026-08"])
        files, checks = build_all(timesheets, conf, {})
        self.assertEqual(sorted(checks), ["2026-09"])
        self.assertFalse(any("2026-08" in name for name in files))
        with tempfile.TemporaryDirectory() as tmp:
            vault = Vault(Path(tmp), MASTER)
            august = vault.month_path("output", "2026-08")
            vault.write_json(august, {"month": "2026-08", "files": {"kept.csv": "ZnJvemVu"}})
            vault.write_json(vault.month_path("output", "2026-07"), {"month": "2026-07", "files": {}})
            write_tree(vault, files, {"2026-08"})
            self.assertEqual(vault.read_json(august)["files"], {"kept.csv": "ZnJvemVu"})
            self.assertEqual(sorted(vault.months("output")), ["2026-08", "2026-09"])
            self.assertEqual(sorted(vault.months("timesheets")), ["2026-09"])


class EncryptedRepoTests(unittest.TestCase):
    """Everything in the private data repo is encrypted, so a copy of it gives nothing away."""

    def test_files_round_trip_and_are_bound_to_their_path(self):
        from etjm_vault import NotEncrypted, Vault

        with tempfile.TemporaryDirectory() as tmp:
            vault = Vault(Path(tmp), MASTER)
            vault.write_json("settings.json", {"employees": [{"name": "Ann Example"}]})
            self.assertEqual(vault.read_json("settings.json"), {"employees": [{"name": "Ann Example"}]})
            text = (Path(tmp) / "settings.json").read_text()
            self.assertNotIn("Ann", text)
            self.assertFalse(vault.write_json("settings.json", {"employees": [{"name": "Ann Example"}]}))
            self.assertEqual((Path(tmp) / "settings.json").read_text(), text)  # unchanged: no new commit

            (Path(tmp) / "users.json").write_text(text)  # moved to another file
            with self.assertRaises(ValueError):
                vault.read("users.json")
            envelope = json.loads(text)
            envelope["data"] = envelope["data"][:-4] + ("AAAA" if not envelope["data"].endswith("AAAA") else "BBBB")
            (Path(tmp) / "settings.json").write_text(json.dumps(envelope))
            with self.assertRaises(ValueError):
                vault.read("settings.json")
            (Path(tmp) / "settings.json").write_text('{"employees": []}')
            with self.assertRaises(NotEncrypted):
                vault.read("settings.json")
            vault.write("settings.json", b"{}")
            with self.assertRaises(ValueError):
                Vault(Path(tmp), bytes(32)).read("settings.json")  # another MASTER_KEY

    def test_a_copy_of_the_repo_names_nobody(self):
        import base64
        import openpyxl
        import subprocess

        from etjm_vault import Vault

        tools = Path(__file__).resolve().parents[1] / "tools"
        keys, _, _ = world()
        with tempfile.TemporaryDirectory() as tmp:
            data, site = Path(tmp) / "data", Path(tmp) / "site"
            site.mkdir()
            vault = Vault(data, MASTER)
            vault.write_json("settings.json", {
                "employees": [{"name": "Ann Example", "rechargeRate": 160}, {"name": "Ben Sample", "rechargeRate": 215}],
                "admins": [{"name": "Dee Office"}]})
            workbook = openpyxl.Workbook()
            workbook.active.append(["Project No.", "Job Description", "Job Location", "Owner", "Project Task No.", "Status"])
            workbook.active.append(["JOB000010", "Secret tank repaint", "SE", 80, "REVEX", "Open"])
            buffer = io.BytesIO()
            workbook.save(buffer)
            vault.write_json("input/report.json", {"name": "Confidential report.xlsx", "sha256": "x", "uploaded": "2026-10-01",
                                                   "xlsx": base64.b64encode(buffer.getvalue()).decode()})
            env = {**os.environ, "MASTER_KEY": base64.b64encode(MASTER).decode()}
            env.pop("GITHUB_TOKEN", None)
            no_issues = Path(tmp) / "none.json"
            no_issues.write_text("[]")
            run = lambda *args: subprocess.run([sys.executable, *args], env=env, capture_output=True, text=True, check=True).stdout  # noqa: E731
            log = run(str(tools / "publish_website.py"), "--data-dir", str(data), "--site-dir", str(site),
                      "--data-repo", "example/data", "--issues-file", str(no_issues))

            state = vault.read_json("users.json")
            saved = sheet(jobs=[{"job": "JOB000010", "hours": {"REVEX": {"3": 7.5}}}])
            issues = Path(tmp) / "issues.json"
            issues.write_text(json.dumps([ts_issue(1, keys, state, "ann.example", saved)]))
            log += run(str(tools / "build_outputs.py"), "--data-dir", str(data), "--issues-file", str(issues))

            secrets = ["Ann Example", "ann.example", "Ben Sample", "ben.sample", "Dee Office", "dee.office",
                       "Secret tank", "Confidential", "JOB000010", "Sep 2026", "2026-09"]
            for path in [*data.rglob("*"), *site.rglob("*")]:
                if path.is_file():
                    text = path.read_bytes().decode("utf-8", "replace") + path.name
                    for secret in secrets:
                        self.assertFalse(secret in text, f"{secret} is readable in {path.relative_to(tmp)}")
            for secret in secrets[:-2]:
                self.assertFalse(secret in log, f"{secret} is in the Actions log")

            (month,) = vault.months("output").values()
            journal = base64.b64decode(month[1]["files"]["Timewriting Journal - Sep 2026.csv"]).decode("utf-8-sig")
            self.assertIn("Ann Example", journal)
            self.assertIn("1200", journal)


class SettingsTests(unittest.TestCase):
    def test_defaults_and_validation(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "settings.json"
            path.write_text(json.dumps({"defaultRechargeRate": 170, "journal": {"balancingGlAccount": "9012"}}))
            loaded = load_settings(path)
            self.assertNotIn("defaultRechargeRate", loaded)
            self.assertEqual(loaded["journal"]["balancingGlAccount"], "9012")
            self.assertEqual(loaded["journal"]["dateFormat"], "%d/%m/%Y")
            self.assertEqual(loaded["employees"], [])
            path.write_text(json.dumps({"employees": [{"name": "Ann Example", "rechargeRate": "lots"}]}))
            with self.assertRaises(ValueError):
                load_settings(path)

    def test_usernames(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "settings.json"
            path.write_text(json.dumps({"employees": [{"name": "Ann Example"}, {"name": "Tāne O'Neil-Smith", "username": "TANE"}]}))
            self.assertEqual([e["username"] for e in load_settings(path)["employees"]], ["ann.example", "tane"])
            path.write_text(json.dumps({"employees": [{"name": "Ann Example"}, {"name": "Ann Other", "username": "ann.example"}]}))
            with self.assertRaises(ValueError):
                load_settings(path)

    def test_admins(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "settings.json"
            path.write_text(json.dumps({"employees": [{"name": "Ann Example"}], "admins": [{"name": "Dee Office"}, {"name": "Ann Example"}]}))
            everyone = {p["username"]: p for p in people(load_settings(path))}
            self.assertEqual(everyone["ann.example"], {"name": "Ann Example", "username": "ann.example", "timesheets": True, "admin": True, "email": None})
            self.assertEqual(everyone["dee.office"], {"name": "Dee Office", "username": "dee.office", "timesheets": False, "admin": True, "email": None})
            path.write_text(json.dumps({"employees": [{"name": "Ann Example"}], "admins": [{"name": "Someone Else", "username": "ann.example"}]}))
            with self.assertRaises(ValueError):
                load_settings(path)

    def test_business_central(self):
        """The Admin page's Business Central settings are kept as they are, and must be texts."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "settings.json"
            bc = {"tenant": "example.co.nz", "batch": "TIME", "capexAccount": "6100"}
            path.write_text(json.dumps({"businessCentral": bc}))
            self.assertEqual(load_settings(path)["businessCentral"], bc)
            self.assertNotIn("businessCentral", load_settings(Path(tmp) / "missing.json"))
            for wrong in ([], {"batch": 7}, {"batch": "x" * 301}):
                path.write_text(json.dumps({"businessCentral": wrong}))
                with self.assertRaises(ValueError):
                    load_settings(path)



class ZeroHourJobTests(unittest.TestCase):
    """Jobs with no hours stay on the timesheet and in the workbook, but not in the journal."""

    def test_kept_round_tripped_and_left_out_of_the_journal(self):
        data = sheet(jobs=[
            {"job": "JOB000010", "location": "SE", "owner": "30", "description": "Tank repaint", "hours": {"REVEX": {"1": 2}}},
            {"job": "JOB000011", "location": "DU", "owner": "80", "description": "No hours yet", "hours": {}},
            {"job": "JOB000012", "description": "Empty CAPEX line", "hours": {"CAPEX": {}}},
            {"job": "JOB000013", "description": "Listed twice", "hours": {}},
            {"job": "JOB000013", "description": "Listed twice", "hours": {"REVEX": {"2": 1}}},
        ])
        ts = parse_timesheet(data)
        self.assertEqual([(line.job, line.task, line.total) for line in ts.jobs],
                         [("JOB000010", "REVEX", 2), ("JOB000011", "", 0), ("JOB000012", "CAPEX", 0), ("JOB000013", "REVEX", 1)])
        again = parse_timesheet(ts.to_json())  # the verified copy keeps them too
        self.assertEqual([(line.job, line.task) for line in again.jobs], [(line.job, line.task) for line in ts.jobs])
        self.assertEqual(ts.to_json()["jobs"][1]["hours"], {})
        ts.username = "ann.example"
        lines, _ = build_journal([ts], settings(), 2026, 9)
        self.assertEqual([line.project for line in lines if line.project], ["JOB000010", "JOB000013"])

    def test_the_workbook_lists_them(self):
        import openpyxl

        timesheets, _ = collected([sheet(jobs=[
            {"job": "JOB000010", "hours": {"REVEX": {"1": 2}}},
            {"job": "JOB000011", "description": "No hours yet", "hours": {}},
            {"job": "JOB000099", "description": "Closed job", "hours": {}},
        ])])
        jobs_index = {"JOB000010": {"no": "JOB000010", "desc": "Tank repaint", "loc": "SE", "owner": "30", "tasks": ["REVEX"]},
                      "JOB000011": {"no": "JOB000011", "desc": "No hours yet", "loc": "DU", "owner": "80", "tasks": ["CAPEX", "REVEX"]}}
        files, checks = build_all(timesheets, settings(), jobs_index)
        name = next(n for n in files if "Engineer timewriting" in n)
        ws = openpyxl.load_workbook(io.BytesIO(files[name])).active
        rows = {r[0]: r for r in ws.iter_rows(values_only=True) if r and str(r[0] or "").startswith("JOB")}
        self.assertEqual(set(rows), {"JOB000010", "JOB000011", "JOB000099"})
        self.assertEqual(rows["JOB000011"][4], "CAPEX/REVEX")  # what it's open for
        self.assertEqual(rows["JOB000011"][5] or "", "")  # nothing wrong with a job that has no hours
        self.assertEqual(rows["JOB000099"][5], "Not an open job")
        self.assertFalse(any("JOB000011" in c or "JOB000099" in c for c in checks["2026-09"]))
        # ...and a second copy without them, for the Admin page's "Include jobs with no hours" off
        short = next(n for n in files if n.endswith("(only jobs with hours).xlsx"))
        self.assertEqual(short, name[:-5] + " (only jobs with hours).xlsx")
        ws = openpyxl.load_workbook(io.BytesIO(files[short])).active
        self.assertEqual([r[0] for r in ws.iter_rows(values_only=True) if r and str(r[0] or "").startswith("JOB")], ["JOB000010"])
        readme = files[next(n for n in files if n.endswith("README.md"))].decode()
        self.assertNotIn("only jobs with hours", readme)



class DeletedMonthTests(unittest.TestCase):
    """An engineer can delete a month that isn't closed; nobody else can, and old copies can't undo it."""

    def mirror_after(self, issues, tags, mirror=None, closed=()):
        found, warnings = collect_timesheets(issues, tags, mirror, set(closed))
        files, checks = build_all(found, settings(closedMonths=list(closed)), {})
        kept = {}
        for name, content in files.items():
            if name.startswith("timesheets/"):
                data = json.loads(content)
                ts = parse_timesheet(data)
                ts.username = data["username"]
                kept[(ts.username, ts.month_key)] = ts
        return found, files, kept, warnings

    def test_delete_then_replay_then_save_again(self):
        keys, everyone, state = world()
        tags = keys_by_tag(keys, state)
        work = sheet(saved="2026-10-01T00:00:00Z", jobs=[{"job": "JOB000010", "hours": {"REVEX": {"1": 2}}}])
        saved = ts_issue(1, keys, state, "ann.example", work)
        _, files, mirror, _ = self.mirror_after([saved], tags)
        self.assertTrue(any(n.startswith("output/2026-09/") for n in files))

        gone = ts_issue(1, keys, state, "ann.example", {"v": 1, "employee": "Ann Example", "month": "2026-09",
                                                         "savedAt": "2026-10-03T00:00:00Z", "deleted": True, "jobs": [], "other": []})
        found, files, mirror, _ = self.mirror_after([gone], tags, mirror)
        self.assertEqual([t.deleted for t in found], [True])
        self.assertFalse(any(n.startswith("output/") for n in files))  # no journal lines, no workbook
        self.assertTrue(mirror[("ann.example", "2026-09")].deleted)  # remembered

        _, files, mirror, _ = self.mirror_after([saved], tags, mirror)  # the older copy put back
        self.assertFalse(any(n.startswith("output/") for n in files))

        again = ts_issue(1, keys, state, "ann.example", dict(work, savedAt="2026-10-04T00:00:00Z"))
        _, files, mirror, _ = self.mirror_after([again], tags, mirror)
        self.assertTrue(any(n.startswith("output/2026-09/") for n in files))
        self.assertFalse(mirror[("ann.example", "2026-09")].deleted)

    def test_nobody_else_can_delete_it_and_closed_months_stay(self):
        keys, everyone, state = world()
        tags = keys_by_tag(keys, state)
        saved = ts_issue(1, keys, state, "ann.example", sheet(saved="2026-10-01T00:00:00Z", jobs=[{"job": "JOB000010", "hours": {"REVEX": {"1": 2}}}]))
        _, _, mirror, _ = self.mirror_after([saved], tags)
        deletion = {"v": 1, "employee": "Ann Example", "month": "2026-09", "savedAt": "2026-10-03T00:00:00Z", "deleted": True, "jobs": [], "other": []}
        forged = ts_issue(2, keys, state, "ben.sample", deletion)  # Ben's key, Ann's title
        forged["title"] = saved["title"]
        found, files, _, warnings = self.mirror_after([forged], tags, mirror)
        self.assertEqual([(t.username, t.deleted) for t in found], [("ann.example", False)])
        self.assertTrue(any(n.startswith("output/2026-09/") for n in files))
        self.assertEqual(len(warnings), 1)
        mine = ts_issue(3, keys, state, "ann.example", deletion)
        found, _, _, _ = self.mirror_after([mine], tags, mirror, closed=["2026-09"])
        self.assertFalse(any(t.deleted for t in found))  # a closed month can't be deleted


if __name__ == "__main__":
    unittest.main()
